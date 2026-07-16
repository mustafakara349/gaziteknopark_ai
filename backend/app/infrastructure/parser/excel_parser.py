# app/infrastructure/parser/excel_parser.py
import os
import logging
from typing import List, Dict, Any
from app.application.interfaces.file_parser import IFileParserStrategy

logger = logging.getLogger("app.infrastructure.parser.excel")

class ExcelParserStrategy(IFileParserStrategy):
    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Dosya bulunamadı: {file_path}")
            
        try:
            import openpyxl
        except ImportError:
            logger.error("openpyxl kütüphanesi bulunamadı! 'pip install openpyxl' çalıştırılmalı.")
            raise RuntimeError("Excel okuma kütüphanesi eksik.")

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            pages = []
            
            for sheet_idx, sheet_name in enumerate(wb.sheetnames):
                sheet = wb[sheet_name]
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    # Boş olmayan satırları al
                    if any(cell is not None for cell in row):
                        row_str = " | ".join([str(cell) if cell is not None else "" for cell in row])
                        rows.append(row_str)
                        
                sheet_content = f"Tablo: {sheet_name}\n" + "\n".join(rows)
                # Her excel sayfasını bir 'sayfa' (page) olarak indeksle
                pages.append({
                    "text": sheet_content,
                    "page_number": sheet_idx + 1
                })
                
            wb.close()
            return pages
        except Exception as e:
            logger.error(f"Excel ayrıştırılırken hata oluştu: {str(e)}")
            raise e
