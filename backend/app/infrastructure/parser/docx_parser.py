# app/infrastructure/parser/docx_parser.py
import os
import logging
from typing import List, Dict, Any
from app.application.interfaces.file_parser import IFileParserStrategy

logger = logging.getLogger("app.infrastructure.parser.docx")

class DocxParserStrategy(IFileParserStrategy):
    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Dosya bulunamadı: {file_path}")
            
        try:
            import docx  # python-docx
        except ImportError:
            logger.error("python-docx kütüphanesi bulunamadı! 'pip install python-docx' çalıştırılmalı.")
            raise RuntimeError("Word okuma kütüphanesi eksik.")

        try:
            doc = docx.Document(file_path)
            full_text = []
            
            # 1. Normal paragrafları oku
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)
                    
            # 2. Tabloları oku
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        full_text.append(" | ".join(row_text))
                    
            # Word belgelerinde sayfa kavramı dinamik olduğu için tek bir sayfa olarak kabul edilir
            text_content = "\n".join(full_text)
            return [{
                "text": text_content,
                "page_number": 1
            }]
        except Exception as e:
            logger.error(f"Docx ayrıştırılırken hata oluştu: {str(e)}")
            raise e
