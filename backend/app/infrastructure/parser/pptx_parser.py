# app/infrastructure/parser/pptx_parser.py
import os
import logging
from typing import List, Dict, Any
from app.application.interfaces.file_parser import IFileParserStrategy

logger = logging.getLogger("app.infrastructure.parser.pptx")

class PptxParserStrategy(IFileParserStrategy):
    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Dosya bulunamadı: {file_path}")
            
        try:
            from pptx import Presentation
        except ImportError:
            logger.error("python-pptx kütüphanesi bulunamadı! 'pip install python-pptx' çalıştırılmalı.")
            raise RuntimeError("PowerPoint okuma kütüphanesi eksik.")

        try:
            prs = Presentation(file_path)
            pages = []
            
            for slide_idx, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text)
                        
                content = "\n".join(slide_text)
                # Her slaytı bir sayfa (page) olarak kabul et
                pages.append({
                    "text": content,
                    "page_number": slide_idx + 1
                })
                
            return pages
        except Exception as e:
            logger.error(f"PowerPoint ayrıştırılırken hata oluştu: {str(e)}")
            raise e
